"""Project 01 — RAG Document Chat: ingestion + retrieval engine.

Pipeline:  document bytes → loader → text splitter → embeddings → vector store → retriever

Supports PDF (``pypdf``), DOCX (``python-docx``), TXT and Markdown.
Everything runs *in-process* so it deploys cleanly on Streamlit Community
Cloud (no local file system, no external vector service required).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

# ---------------------------------------------------------------------------
# Embeddings — fastembed (ONNX runtime) gives real, high-quality vectors with a
# fraction of the footprint of sentence-transformers/torch, which makes it a
# much better fit for Streamlit Community Cloud's ~1 GB containers.
# We degrade gracefully to a deterministic embedder so the demo never crashes.
# ---------------------------------------------------------------------------
_EMBEDDER = None
_EMBEDDER_NAME: str | None = None
DEFAULT_EMBED_MODEL = "BAAI/bge-small-en-v1.5"  # 384-dim, fast + strong on retrieval


class FastEmbedEmbeddings:
    """Minimal LangChain-``Embeddings``-compatible wrapper around fastembed."""

    def __init__(self, model_name: str = DEFAULT_EMBED_MODEL):
        from fastembed import TextEmbedding

        self.model_name = model_name
        self._model = TextEmbedding(model_name=model_name)

    def embed_documents(self, texts):
        return [list(map(float, v)) for v in self._model.embed(list(texts))]

    def embed_query(self, text):
        return list(map(float, next(iter(self._model.query_embed([text])))))


def get_embeddings():
    """Return a cached embeddings object (fastembed primary, deterministic fallback)."""
    global _EMBEDDER, _EMBEDDER_NAME
    if _EMBEDDER is not None:
        return _EMBEDDER
    try:
        from langchain_core.embeddings import Embeddings as _LCEmbeddings

        class _FastEmbed(_LCEmbeddings):
            def __init__(self, model_name: str = DEFAULT_EMBED_MODEL):
                self._impl = FastEmbedEmbeddings(model_name)

            def embed_documents(self, texts):
                return self._impl.embed_documents(texts)

            def embed_query(self, text):
                return self._impl.embed_query(text)

        _EMBEDDER = _FastEmbed()
        _EMBEDDER_NAME = f"{DEFAULT_EMBED_MODEL} (fastembed/ONNX)"
    except Exception:
        from langchain_core.embeddings import DeterministicFakeEmbedding

        _EMBEDDER = DeterministicFakeEmbedding(size=384)
        _EMBEDDER_NAME = (
            "DeterministicFakeEmbedding (fallback — install fastembed for real vectors)"
        )
    return _EMBEDDER


def embedder_name() -> str:
    get_embeddings()
    return _EMBEDDER_NAME or "unknown"


def get_vector_store(embeddings=None, metadatas: list[dict] | None = None, texts: list[str] | None = None):
    """Build a FAISS vector store, falling back to an in-memory store."""
    embeddings = embeddings or get_embeddings()
    texts = texts or []
    metadatas = metadatas or [{} for _ in texts]
    try:
        from langchain_community.vectorstores import FAISS

        return FAISS.from_texts(texts, embeddings, metadatas=metadatas)
    except Exception:
        from langchain_core.vectorstores import InMemoryVectorStore

        store = InMemoryVectorStore(embedding=embeddings)
        if texts:
            store.add_texts(texts, metadatas=metadatas)
        return store


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            pages.append(f"[page {i}]\n{page_text}")
    return _clean("\n\n".join(pages))


def load_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _clean("\n".join(parts))


def load_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return _clean(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return _clean(data.decode("utf-8", errors="ignore"))


def load_document(filename: str, data: bytes) -> str:
    """Route a file to the right loader based on its extension."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return load_pdf(data)
    if name.endswith(".docx"):
        return load_docx(data)
    if name.endswith((".txt", ".md", ".csv", ".json")):
        return load_text(data)
    # best effort
    text = load_text(data)
    if not text:
        raise ValueError(f"Unsupported or empty file: {filename}")
    return text


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------
def split_text(text: str, source: str, chunk_size: int = 1000, chunk_overlap: int = 150):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        add_start_index=True,
    )
    docs = splitter.create_documents([text], metadatas=[{"source": source}])
    # Drop empty chunks and assign a stable chunk id.
    out: list[Document] = []
    for idx, d in enumerate(docs):
        if d.page_content.strip():
            d.metadata["chunk"] = idx
            out.append(d)
    return out


@dataclass
class IngestResult:
    source: str
    n_chars: int
    chunks: list[Document] = field(default_factory=list)

    @property
    def n_chunks(self) -> int:
        return len(self.chunks)


def ingest(filename: str, data: bytes, chunk_size: int = 1000, chunk_overlap: int = 150) -> IngestResult:
    """Full load → split step for a single uploaded document."""
    text = load_document(filename, data)
    chunks = split_text(text, filename, chunk_size, chunk_overlap)
    return IngestResult(source=filename, n_chars=len(text), chunks=chunks)


def build_retriever(chunks: Iterable[Document], k: int = 4):
    """Build a retriever from chunk documents."""
    chunks = list(chunks)
    if not chunks:
        raise ValueError("No chunks to index — upload a document first.")
    store = get_vector_store(
        texts=[c.page_content for c in chunks],
        metadatas=[c.metadata for c in chunks],
    )
    return store.as_retriever(search_kwargs={"k": min(k, len(chunks))})


def format_sources(docs: list[Document]) -> list[dict]:
    """Turn retrieved documents into compact, display-friendly source records."""
    seen: dict[tuple, dict] = {}
    for rank, d in enumerate(docs, start=1):
        md = d.metadata or {}
        key = (md.get("source"), md.get("chunk"))
        snippet = " ".join(d.page_content.split())[:280]
        if key in seen:
            seen[key]["score_rank"] = min(seen[key]["score_rank"], rank)
            continue
        seen[key] = {
            "source": md.get("source", "unknown"),
            "chunk": md.get("chunk"),
            "page": md.get("page"),
            "rank": rank,
            "snippet": snippet,
        }
    return sorted(seen.values(), key=lambda s: s["score_rank"] if "score_rank" in s else s["rank"])
